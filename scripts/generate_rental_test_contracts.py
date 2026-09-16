from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
WORK_DIR = ROOT / "tmp" / "rental-contracts"

FONT = "Arial"
INK = RGBColor(28, 36, 45)
NAVY = RGBColor(31, 58, 95)
MUTED = RGBColor(92, 101, 112)
LIGHT = "E9EEF4"


CONTRACTS = [
    {
        "slug": "01_contract_tel_aviv",
        "title": "חוזה שכירות למגורים - תל אביב",
        "subtitle": "דירה ברחוב דיזנגוף 118, דירה 7, תל אביב-יפו",
        "date": "26 באוגוסט 2026",
        "sections": [
            ("1. הצדדים והדירה", [
                "בין מר יונתן ברק, תעודת זהות 900000011, שמענו לצורך הסכם זה ברחוב בן יהודה 64, תל אביב-יפו (להלן: \"המשכיר\"), לבין גב' נועה מזרחי, תעודת זהות 900000029, שמענה לצורך הסכם זה בדירה המושכרת (להלן: \"השוכרת\").",
                "המשכיר משכיר לשוכרת דירת מגורים בת שלושה חדרים בקומה השלישית, ברחוב דיזנגוף 118, דירה 7, תל אביב-יפו, לרבות מחסן מס' 7 ושימוש רגיל ברכוש המשותף (להלן: \"הדירה\").",
            ]),
            ("2. הצהרות ומצב הדירה", [
                "המשכיר מצהיר כי הוא בעל הזכות להשכיר את הדירה וכי אין זכות נוגדת של צד שלישי. השוכרת מאשרת שביקרה בדירה ומצאה אותה מתאימה לצרכיה במצבה הנוכחי, ומוותרת באופן בלתי חוזר על כל טענה או דרישה בקשר למצב הדירה, בין אם הייתה ידועה לה במועד החתימה ובין אם לאו.",
                "בדירה נשארים מקרר, תנור אפייה, שני מזגנים, ארון קיר בחדר השינה ושולחן אוכל. ידועים סימני רטיבות ישנים בתקרת חדר הרחצה וסדק בקיר הסלון.",
            ]),
            ("3. מטרת השכירות", [
                "השימוש בדירה יהיה למגורי השוכרת ובני משפחתה בלבד. אין לעשות בדירה שימוש עסקי, להשכיר אותה בשכירות משנה או להעביר את החזקה בה לאחר ללא הסכמה מראש ובכתב של המשכיר.",
            ]),
            ("4. תקופת השכירות", [
                "תקופת השכירות היא שנים עשר חודשים, מיום 1 בספטמבר 2026 ועד יום 31 באוגוסט 2027. החזקה תימסר ביום תחילת התקופה כשהדירה פנויה מכל אדם וחפץ שאינו מפורט בהסכם.",
            ]),
            ("5. דמי השכירות", [
                "דמי השכירות החודשיים הם 6,000 ש\"ח. השוכרת תשלם אותם עד ליום הראשון בכל חודש בהעברה בנקאית לחשבון שפרטיו יימסרו לה בכתב. איחור של יותר משבעה ימים ייחשב הפרה יסודית.",
            ]),
            ("6. תשלומים נוספים", [
                "בנוסף לדמי השכירות תישא השוכרת בארנונה, מים, חשמל, גז ותשלומי ועד הבית השוטפים. השוכרת תישא גם בדמי ביטוח המבנה ובכל חיוב מיוחד של נציגות הבית המשותף, לרבות החלפת המעלית, איטום הגג והשבחת הלובי, גם אם החיוב הוטל על בעלי הדירות.",
            ]),
            ("7. תיקונים ותחזוקה", [
                "השוכרת תתקן על חשבונה כל נזק שנגרם עקב שימוש בלתי סביר שלה. המשכיר יתקן תקלה אחרת שאינה קלת ערך בתוך 45 ימים ממועד הודעת השוכרת. ליקוי דחוף שאינו מאפשר מגורים סבירים בדירה יתוקן בתוך שבעה ימים ממועד קבלת הודעת השוכרת.",
                "אם השוכרת הזמינה בעל מקצוע ללא אישור מוקדם ובכתב של המשכיר, היא לא תהיה זכאית להחזר הוצאה כלשהי, גם במקרה דחוף.",
            ]),
            ("8. שינויים בדירה", [
                "השוכרת לא תבצע שינוי או התקנה קבועה בלי הסכמה מראש ובכתב של המשכיר. שינוי שאושר יישאר בדירה ללא תמורה, אלא אם המשכיר דרש להחזיר את המצב לקדמותו.",
            ]),
            ("9. בטוחות", [
                "להבטחת התחייבויות השוכרת תמסור השוכרת למשכיר ערבות בנקאית בסך 30,000 ש\"ח לפני קבלת החזקה בדירה.",
                "המשכיר רשאי לממש את הערובה בשל כל הפרה, לפי שיקול דעתו, ללא הודעה מוקדמת או אפשרות לתקן את ההפרה.",
            ]),
            ("10. כניסת המשכיר לדירה", [
                "המשכיר או מי מטעמו רשאים להיכנס לדירה בשעות היום לבדיקה, לתיקון או להצגתה לאחרים. השוכרת מסכימה לכך מראש, ללא צורך בתיאום או בהודעה מוקדמת.",
            ]),
            ("11. בעלי חיים", [
                "אסור להחזיק בדירה כלב, חתול או בעל חיים אחר, גם באופן זמני, ללא הסכמה פרטנית מראש ובכתב של המשכיר.",
            ]),
            ("12. ביטול מוקדם", [
                "המשכיר רשאי להביא את החוזה לסיום מכל סיבה בהודעה של 30 ימים מראש. לשוכרת לא תהיה זכות מקבילה לבטל את החוזה לפני תום התקופה, והיא תישא בדמי השכירות עד סופה גם אם תפנה את הדירה קודם לכן.",
            ]),
            ("13. הפרה ופינוי", [
                "אי תשלום דמי שכירות, ביצוע שימוש אסור בדירה או אי פינוי במועד הם הפרות יסודיות. במקרה של הפרה יסודית שלא תוקנה בתוך שבעה ימים מדרישה בכתב, יהיה המשכיר רשאי לבטל את ההסכם ולנקוט בכל סעד המותר לפי דין.",
            ]),
            ("14. החזרת הדירה והערובה", [
                "בתום השכירות תחזיר השוכרת את הדירה כשהיא פנויה ונקייה ובמצב שבו נמסרה, למעט בלאי סביר. הערובה או יתרתה יוחזרו לשוכרת בתוך 90 ימים ממועד פינוי הדירה ולאחר הצגת אישורים על סילוק כל החובות, לפי המאוחר.",
            ]),
            ("15. העברת זכויות המשכיר", [
                "המשכיר רשאי להעביר את זכויותיו בדירה לאחר, ובלבד שזכויות השוכרת לפי הסכם זה יישמרו. הודעה על זהות הרוכש ודרכי ההתקשרות עמו תימסר לשוכרת בכתב.",
            ]),
            ("16. הודעות", [
                "הודעה לפי הסכם זה תימסר בדואר אלקטרוני או במסרון למספרים שימסרו הצדדים, ותיחשב כמסורה ביום העסקים הבא. שינוי בהסכם יהיה בכתב ובחתימת שני הצדדים בלבד.",
            ]),
            ("17. הדין החל", [
                "על הסכם זה יחולו דיני מדינת ישראל. סמכות השיפוט המקומית תהיה לבית המשפט המוסמך בתל אביב-יפו.",
            ]),
        ],
    },
    {
        "slug": "02_contract_jerusalem_furnished",
        "title": "הסכם שכירות לדירה מרוהטת - ירושלים",
        "subtitle": "דירה ברחוב עזה 43, דירה 2, ירושלים",
        "date": "26 באוגוסט 2026",
        "sections": [
            ("1. מבוא והצדדים", [
                "בין גב' רונית שלו, תעודת זהות 900000037, שמענה ברחוב הפלמ\"ח 19, ירושלים (להלן: \"בעלת הדירה\"), לבין מר אדם לוי, תעודת זהות 900000045, שמענו בתקופת השכירות בדירה המושכרת (להלן: \"השוכר\").",
                "בעלת הדירה משכירה לשוכר דירה מרוהטת בת שני חדרים בקומה הראשונה, ברחוב עזה 43, דירה 2, ירושלים, ללא חניה וללא מחסן (להלן: \"הדירה\").",
            ]),
            ("2. תכולה ומצב הדירה", [
                "בדירה נשארים מיטה זוגית, ארון בגדים, ספה, שולחן עבודה, מכונת כביסה, מקרר, כיריים ושני מזגנים. השוכר מצהיר שבדק את הדירה והיא מתאימה למגוריו.",
                "בעלת הדירה מודיעה כי קיימת שריטה בפרקט בחדר השינה וכי צפויים רעשי בנייה בשעות היום מהמגרש הסמוך. אין מצורף להסכם פרוטוקול מסירה או תיעוד מצולם של מצב הדירה.",
            ]),
            ("3. מטרת השכירות ושימוש", [
                "הדירה תשמש למגורי השוכר בלבד. השוכר לא יעביר את זכויותיו, לא ישכיר בשכירות משנה ולא יארח אדם דרך קבע ללא אישור מראש ובכתב של בעלת הדירה.",
            ]),
            ("4. תקופת השכירות", [
                "תקופת השכירות היא עשרה חודשים, מיום 1 באוקטובר 2026 ועד יום 31 ביולי 2027. אין לשוכר זכות לסיום מוקדם ואין מנגנון להצעת שוכר חלופי.",
            ]),
            ("5. דמי השכירות", [
                "דמי השכירות החודשיים הם 4,800 ש\"ח וישולמו עד ליום החמישי בכל חודש בהוראת קבע. השוכר ימסור במעמד החתימה המחאה נוספת בסך חודש שכירות להבטחת התשלום האחרון.",
            ]),
            ("6. הארכת השכירות", [
                "לשוכר אפשרות לבקש הארכה של עשרה חודשים נוספים, אך הארכה תיעשה רק בהסכמה חדשה ובכתב של הצדדים. אם בעלת הדירה תסכים, העלאת דמי השכירות בתקופה הנוספת תהיה 8% לעומת דמי השכירות האחרונים. לא נקבע מועד למתן הודעה על הרצון בהארכה.",
            ]),
            ("7. תשלומים והוצאות", [
                "השוכר יישא בארנונה, מים, חשמל, גז, אינטרנט ותשלומי ועד הבית השוטפים. נוסף על כך ישלם השוכר את עמלת המתווך שפעל מטעם בעלת הדירה בסך 4,800 ש\"ח ואת חלק הדירה בהחלפה המתוכננת של משאבת המים המרכזית בבניין.",
            ]),
            ("8. תיקונים ובלאי", [
                "השוכר אחראי לכל תיקון בדירה שמחירו עד 1,000 ש\"ח, גם אם מקור התקלה בבלאי סביר או במערכת קבועה. בעלת הדירה תתקן תקלה אחרת שאינה באחריות השוכר בתוך 30 ימים, וליקוי דחוף שאינו מאפשר מגורים סבירים בתוך שלושה ימים ממועד הודעה.",
                "השוכר מוותר באופן בלתי חוזר על כל טענה, דרישה או זכות לפיצוי בשל תקלה, הפרעה, פגם או אי התאמה בדירה במהלך תקופת השכירות.",
            ]),
            ("9. בטוחות", [
                "השוכר יפקיד פיקדון כספי בסך 9,600 ש\"ח וימציא ערב אישי להבטחת התחייבויותיו. בעלת הדירה תהיה רשאית לממש את הפיקדון בגין דמי שכירות שלא שולמו, תשלומים שוטפים שלא שולמו, נזק שנגרם בשימוש בלתי סביר או אי פינוי, לאחר התראה של שבעה ימים.",
            ]),
            ("10. השבת הפיקדון", [
                "הפיקדון או יתרתו יוחזרו לשוכר בתוך 75 ימים ממועד החזרת החזקה בדירה, ורק לאחר המצאת אישורי סגירת חשבונות. בעלת הדירה רשאית לקזז מהפיקדון גם עלות ניקיון מקצועי קבועה בסך 1,200 ש\"ח ללא צורך בהצגת חשבונית.",
            ]),
            ("11. בעלי חיים", [
                "אין להחזיק בדירה כלב, חתול או בעל חיים אחר. הפרת הוראה זו תיחשב הפרה יסודית גם אם לא נגרם נזק או מטרד.",
            ]),
            ("12. כניסת בעלת הדירה", [
                "בעלת הדירה רשאית להיכנס לדירה למטרה סבירה, לאחר הודעה מראש של 12 שעות ובשעות מקובלות, לצורך תיקון, בדיקה או הצגת הדירה. במקרה חירום מותרת כניסה ללא הודעה.",
            ]),
            ("13. ביטול ללא עילה", [
                "כל אחד מהצדדים רשאי לבטל את ההסכם ללא הפרה: בעלת הדירה בהודעה של 90 ימים מראש והשוכר בהודעה של 60 ימים מראש. הביטול יהיה בכתב.",
            ]),
            ("14. הפרות וסעדים", [
                "איחור של עשרה ימים בתשלום, גרימת נזק מהותי, שימוש שאינו למגורים או אי פינוי במועד יהוו הפרה יסודית. הצד הנפגע רשאי לבטל את ההסכם לאחר שנתן לצד המפר שבעה ימים לתקן הפרה הניתנת לתיקון.",
            ]),
            ("15. החזרת הדירה", [
                "בתום השכירות יחזיר השוכר את הדירה פנויה, נקייה, מסוידת בצבע לבן ועם מלוא התכולה. השוכר יישא בעלות כל חוסר או נזק לתכולה, ללא הפחתה בשל גיל הפריט או בלאי סביר.",
            ]),
            ("16. הודעות וכללי", [
                "הודעות יישלחו בכתב בדואר אלקטרוני או במסרון. הימנעות של צד ממימוש זכות לא תיחשב ויתור. שינוי להסכם יהיה תקף רק אם נערך בכתב ונחתם על ידי שני הצדדים.",
            ]),
            ("17. דין וסמכות שיפוט", [
                "על ההסכם יחולו דיני מדינת ישראל וסמכות השיפוט המקומית תהיה לבית המשפט המוסמך בירושלים.",
            ]),
        ],
    },
    {
        "slug": "03_contract_haifa_compliant",
        "title": "חוזה שכירות למגורים - חיפה",
        "subtitle": "דירה ברחוב הנביאים 12, דירה 5, חיפה",
        "date": "26 באוגוסט 2026",
        "sections": [
            ("1. הצדדים והתקשרות", [
                "בין גב' מיכל ארז, תעודת זהות 900000052, שמענה ברחוב יפה נוף 31, חיפה (להלן: \"המשכירה\"), לבין מר תומר גפן, תעודת זהות 900000060, שמענו בתקופת השכירות בדירה המושכרת (להלן: \"השוכר\").",
                "המשכירה מצהירה כי היא בעלת הזכות להשכיר את הדירה וכי זכויות השוכר לפי חוזה זה יהיו חופשיות מזכות נוגדת של צד שלישי. ההסכם נערך בכתב, וכל צד יקבל עותק חתום.",
            ]),
            ("2. תיאור הדירה והתכולה", [
                "הדירה נמצאת ברחוב הנביאים 12, דירה 5, חיפה, בקומה השנייה, ובה שלושה חדרים, מטבח, חדר רחצה ומרפסת שירות. לדירה מוצמדים חניה מס' 5 ומחסן מס' 5.",
                "בדירה נשארים מקרר, תנור, כיריים, מכונת כביסה, שלושה מזגנים, ארון קיר ושולחן אוכל. רשימת ציוד מלאה ומצב כל פריט יתועדו בפרוטוקול המסירה המצורף כנספח א'.",
            ]),
            ("3. מצב הדירה וראויות למגורים", [
                "המשכירה תמסור את הדירה כשהיא פנויה, נקייה ומתאימה למוסכם, עם מערכת ניקוז וסילוק שפכים תקינה, אספקת מי שתייה, מערכות חשמל ותאורה, פתחי אוורור ותאורה טבעית, דלתות וחלונות הניתנים לסגירה, דלת כניסה הניתנת לנעילה ומחיצה בין השירותים לשאר הדירה, וללא סיכון בלתי סביר לבריאות או לבטיחות.",
                "המשכירה מצהירה כי לא ידוע לה על פגם שאינו קל ערך או על הפרעה מהותית בדירה או בסביבתה, למעט סדק שטחי באורך כ-15 ס\"מ מעל חלון הסלון שאינו פעיל ואינו מעיד על רטיבות. הסדק יתועד בתמונות בפרוטוקול המסירה.",
            ]),
            ("4. מטרת השכירות", [
                "הדירה תשמש למגורי השוכר ובת זוגו בלבד. עבודה משרדית שקטה מהבית מותרת ובלבד שאין קבלת קהל, שילוט, מטרד או שינוי ייעוד. השכירות אינה מוגנת לפי חוק הגנת הדייר ולא שולמו דמי מפתח.",
            ]),
            ("5. תקופת השכירות ומסירה", [
                "תקופת השכירות היא שנים עשר חודשים, מיום 1 בספטמבר 2026 ועד יום 31 באוגוסט 2027. החזקה והמפתחות יימסרו ביום 1 בספטמבר 2026 לאחר חתימת פרוטוקול מסירה הכולל את מצב הדירה, הציוד, הליקויים הקיימים, קריאות המונים ותמונות מוסכמות.",
            ]),
            ("6. דמי השכירות ואופן התשלום", [
                "דמי השכירות החודשיים הם 4,900 ש\"ח. השוכר ישלם אותם עד ליום הראשון בכל חודש בהעברה בנקאית. תשלום ייחשב שבוצע במועד שבו זוכה חשבון המשכירה, אלא אם העיכוב נגרם מסיבה התלויה במשכירה.",
            ]),
            ("7. אופציה להארכה", [
                "לשוכר אופציה, ודמי השכירות בתקופה הנוספת יהיו 5,047 ש\"ח לחודש. האופציה היא להארכת השכירות ב-12 חודשים, מיום 1 בספטמבר 2027 ועד יום 31 באוגוסט 2028, בכפוף לכך שלא קיימת הפרה יסודית שלא תוקנה.",
                "השוכר יודיע על מימוש האופציה בכתב לפחות 60 ימים לפני תום התקופה הראשונה. דמי השכירות האמורים משקפים העלאה של 3%, ושאר תנאי ההסכם ימשיכו לחול. המשכירה תזכיר לשוכר את תנאי האופציה לפחות 90 ימים לפני תום התקופה הראשונה.",
            ]),
            ("8. תשלומים שוטפים ותשלומי בעלים", [
                "השוכר יישא בתקופת השכירות בארנונה החלה על המחזיק, במים, בחשמל, בגז ובתשלומי ועד הבית המיועדים לאחזקה שוטפת. השוכר יעביר את החשבונות על שמו וישלם אותם במועד.",
                "המשכירה תישא בביטוח המבנה, בהיטלים ובתשלומים החלים על בעלים, בהשבחות וברכישה או החלפה של מתקנים קבועים, בתשלום מיוחד של ועד הבית שאינו אחזקה שוטפת ובדמי תיווך של מתווך שפעל מטעמה. ביטוח תכולה וצד שלישי נתון לבחירת השוכר ועל חשבונו.",
            ]),
            ("9. תיקונים, בלאי ותרופות", [
                "השוכר יתקן על חשבונו ליקוי שנגרם עקב שימוש בלתי סביר שלו או של מי מטעמו. השוכר לא יישא בעלות בלאי סביר או ליקוי שלא נגרם עקב שימוש בלתי סביר שלו.",
                "המשכירה תתקן ליקוי דחוף בתוך 3 ימים. לעניין זה, ליקוי דחוף הוא ליקוי שאינו מאפשר מגורים סבירים. ליקוי אחר שאינו קל ערך יתוקן בתוך זמן סביר ולא יאוחר מ-30 ימים ממועד קבלת הודעה.",
                "אם המשכירה לא תיקנה במועד, רשאי השוכר לפעול לפי התרופות המוקנות לו בדין, לרבות תיקון סביר ודרישת החזר הוצאות או הפחתת דמי שכירות, לאחר מתן הזדמנות מתאימה למשכירה, למעט מקרה דחוף שאינו סובל דיחוי.",
            ]),
            ("10. פרטיות וכניסת המשכירה", [
                "כניסת המשכירה לדירה תהיה רק לאחר תיאום והודעה מראש. ההודעה תינתן 48 שעות מראש, והכניסה תהיה למטרה סבירה, בשעות מקובלות ובתדירות סבירה, לצורך בדיקה, תיקון או הצגה. במקרה חירום ניתן להיכנס מיד, במידה הנדרשת בלבד.",
            ]),
            ("11. בעלי חיים", [
                "השוכר רשאי להחזיק בדירה חתול בית אחד, ובלבד שישמור על ניקיון, ימנע מטרד ויישא בעלות נזק שנגרם בפועל עקב החזקת בעל החיים, למעט בלאי סביר.",
            ]),
            ("12. שינויים בדירה", [
                "השוכר לא יבצע שינוי קבוע בדירה ללא הסכמה מראש ובכתב של המשכירה. הסכמה תפרט אם בתום התקופה יישאר השינוי בדירה או שהשוכר יחזיר את המצב לקדמותו. תיקון חורים קטנים מתליית תמונות ייעשה באופן מקצועי בעת הפינוי.",
            ]),
            ("13. בטוחות ומימושן", [
                "להבטחת התחייבויות השוכר יופקד פיקדון כספי בסך 9,800 ש\"ח. אין בהסכם דרישה לערב אישי או לערבות נוספת הכרוכה בהוצאה כספית לשוכר.",
                "המשכירה רשאית לממש מהפיקדון רק סכום מוכח בגין דמי שכירות שלא שולמו, תשלום שוטף שלא שולם, עלות תיקון נזק שבאחריות השוכר או פיצוי בשל אי פינוי, ולאחר הודעה לשוכר זמן סביר מראש ומתן אפשרות לתקן את הטעון תיקון בתוך זמן סביר.",
                "תוחזר הערובה לשוכר בתוך 30 ימים ממועד השבת הדירה והסדרת החובות המותרים למימוש, לפי המאוחר. ההחזר יכלול את יתרת הפיקדון ופירותיו, בצירוף פירוט בכתב של כל קיזוז.",
            ]),
            ("14. סיום מוקדם ושוכר חלופי", [
                "השוכר רשאי לסיים את השכירות מוקדם בהודעה מראש של 60 ימים, או להציע שוכר חלופי סביר שיקבל על עצמו את יתרת ההסכם וימציא בטוחות חלופיות. המשכירה לא תסרב לשוכר חלופי מטעמים בלתי סבירים ותשיב תשובה מנומקת בתוך שבעה ימים מקבלת מלוא פרטיו.",
                "למשכירה אין זכות לבטל את החוזה ללא עילה. ביטול בשל הפרה ייעשה בהתאם לדין ולאחר מתן ארכה סבירה לתיקון הפרה הניתנת לתיקון.",
            ]),
            ("15. הפרות וסעדים", [
                "איחור העולה על שבעה ימים בתשלום דמי השכירות, שימוש אסור מהותי בדירה, גרימת נזק מהותי בשימוש בלתי סביר או אי פינוי במועד עשויים להיחשב הפרה יסודית בהתאם לנסיבות. אין בהסכם כדי לגרוע מסעד או הגנה קוגנטיים של מי מהצדדים.",
            ]),
            ("16. החזרת הדירה", [
                "בתום השכירות יחזיר השוכר את הדירה כשהיא פנויה מחפציו, נקייה ובמצב שבו נמסרה לפי פרוטוקול המסירה, למעט בלאי סביר. הצדדים יערכו בדיקה משותפת, ירשמו קריאות מונים ויחתמו על פרוטוקול החזרה.",
            ]),
            ("17. העברת זכויות המשכירה", [
                "המשכירה רשאית להעביר את זכויותיה בדירה לאחר, ובלבד שכל זכויות השוכר לפי ההסכם והדין יישמרו. לפני מסירת הדירה לרוכש תודיע המשכירה לשוכר בכתב ותפרט את זהות הרוכש ודרכי ההתקשרות עמו.",
            ]),
            ("18. הודעות ושינויים", [
                "הודעות יישלחו לכתובות הדואר האלקטרוני ולמספרי הטלפון שמסרו הצדדים. הודעה מהותית על הפרה, מימוש אופציה, ביטול או מימוש בטוחה תימסר בכתב. שינוי להסכם יהיה תקף רק אם נעשה בכתב ונחתם בידי שני הצדדים.",
            ]),
            ("19. דין גובר וסמכות שיפוט", [
                "על ההסכם יחולו דיני מדינת ישראל. הוראה שאין להתנות עליה או שניתן להתנות עליה רק לטובת השוכר תגבר על נוסח סותר. סמכות השיפוט המקומית תהיה לבית המשפט המוסמך במחוז חיפה, בכפוף לכל דין.",
            ]),
        ],
    },
]


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent_dxa: int = 120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)


def add_bidi(paragraph, rtl: bool = True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")


def set_run(run, *, size=10.5, bold=False, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "1")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:bidi"), "he-IL")


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = INK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, before, after, color in (
        ("Heading 1", 12.75, 10, 4, NAVY),
        ("Heading 2", 11.5, 8, 4, NAVY),
        ("Heading 3", 10.5, 6, 3, NAVY),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_bidi(paragraph, False)
    prefix = paragraph.add_run("עמוד ")
    set_run(prefix, size=8.5, color=MUTED)
    field_run = paragraph.add_run()
    set_run(field_run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instruction, separate, cached, end])


def add_header_footer(section, title: str):
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    add_bidi(header_p)
    run = header_p.add_run(f"{title}  |  מסמך בדיקה סינתטי")
    set_run(run, size=8.5, bold=True, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.paragraph_format.space_before = Pt(0)
    add_page_number(footer_p)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_signatures(doc: Document):
    heading = doc.add_paragraph("חתימות", style="Heading 1")
    add_bidi(heading)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [4680, 4680], 120)
    labels = (
        ("המשכיר/ה", "השוכר/ת"),
        ("שם, חתימה ותאריך: ____________________", "שם, חתימה ותאריך: ____________________"),
    )
    for row_idx, values in enumerate(labels):
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                shade_cell(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_bidi(p)
            run = p.add_run(value)
            set_run(run, size=9.5, bold=row_idx == 0)


def add_test_notice(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    add_bidi(p)
    run = p.add_run("כל השמות, המספרים והכתובות במסמך זה בדויים ונועדו לבדיקת תוכנה בלבד. אין לחתום עליו ואין להשתמש בו כייעוץ משפטי.")
    set_run(run, size=9, bold=True, color=RGBColor(135, 72, 32))


def build_contract(spec: dict) -> Path:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    # Named override: Israeli legal A4 form factor while retaining the selected
    # standard_business_brief type and spacing system.
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    add_header_footer(section, spec["title"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    add_bidi(p)
    run = p.add_run(spec["title"])
    set_run(run, size=21, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    add_bidi(p)
    run = p.add_run(spec["subtitle"])
    set_run(run, size=11.5, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    add_bidi(p)
    run = p.add_run(f"נערך ונחתם ביום {spec['date']}")
    set_run(run, size=9.5, color=MUTED)
    add_test_notice(doc)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    intro.paragraph_format.space_after = Pt(7)
    add_bidi(intro)
    run = intro.add_run("הואיל והצדדים מבקשים להסדיר שכירות למגורים שאינה מוגנת; לפיכך הוצהר, הוסכם והותנה ביניהם כדלקמן:")
    set_run(run, size=10.5, bold=True)

    for heading_text, paragraphs in spec["sections"]:
        heading = doc.add_paragraph(heading_text, style="Heading 1")
        add_bidi(heading)
        for existing in heading.runs:
            set_run(existing, size=13, bold=True, color=NAVY)
        for body_text in paragraphs:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.widow_control = True
            add_bidi(p)
            run = p.add_run(body_text)
            set_run(run, size=10.25)

    add_signatures(doc)

    props = doc.core_properties
    props.title = spec["title"]
    props.subject = "חוזה שכירות סינתטי לבדיקת אפליקציית RightRent"
    props.author = "RightRent - test fixture generator"
    props.keywords = "שכירות, חוזה, בדיקות תוכנה, ישראל"

    WORK_DIR.mkdir(parents=True, exist_ok=True)
    output = WORK_DIR / f"{spec['slug']}.docx"
    doc.save(output)
    return output


def main():
    for spec in CONTRACTS:
        output = build_contract(spec)
        print(output)


if __name__ == "__main__":
    main()
